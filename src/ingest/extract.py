"""Extract text from uploaded files (PDF, DOCX, TXT, MD)."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pymupdf as fitz
from docx import Document as DocxDocument

from src.models import FileType

MIN_TEXT_CHARS = 20


class ExtractError(Exception):
    """User-visible extraction failure."""


@dataclass(frozen=True)
class PageText:
    """One page (or block) of extracted text."""

    page: int | None  # 1-based for PDF; None for DOCX/TXT/MD
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    file_type: FileType
    pages: list[PageText]
    page_count: int | None


def detect_file_type(path: Path) -> FileType:
    ext = path.suffix.lower().lstrip(".")
    if ext == "pdf":
        return "pdf"
    if ext == "docx":
        return "docx"
    if ext == "txt":
        return "txt"
    if ext == "md":
        return "md"
    raise ExtractError(f"Unsupported file type: .{ext or '(none)'}")


def extract_file(path: Path) -> ExtractedDocument:
    """Dispatch by extension."""
    path = Path(path)
    if not path.is_file():
        raise ExtractError(f"File not found: {path}")

    file_type = detect_file_type(path)
    if file_type == "pdf":
        return extract_pdf(path)
    if file_type == "docx":
        return extract_docx(path)
    if file_type in ("txt", "md"):
        return extract_text_file(path, file_type=file_type)
    raise ExtractError(f"Unsupported file type: {file_type}")


def extract_pdf(path: Path) -> ExtractedDocument:
    """Extract per-page text with 1-based page numbers."""
    try:
        doc = fitz.open(path)
    except Exception as exc:  # noqa: BLE001 — surface as ExtractError
        raise ExtractError(f"Could not open PDF: {exc}") from exc

    page_count = doc.page_count
    pages: list[PageText] = []
    try:
        for i in range(page_count):
            page = doc.load_page(i)
            text = page.get_text("text") or ""
            if text.strip():
                pages.append(PageText(page=i + 1, text=text))
    finally:
        doc.close()

    if not pages:
        raise ExtractError(
            "No extractable text. If this is a scanned PDF, "
            "OCR is not supported in the MVP."
        )

    return ExtractedDocument(
        file_type="pdf",
        pages=pages,
        page_count=page_count,
    )


def extract_docx(path: Path) -> ExtractedDocument:
    try:
        document = DocxDocument(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ExtractError(f"Could not open DOCX: {exc}") from exc

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    text = "\n\n".join(paragraphs)
    if len(text.strip()) < MIN_TEXT_CHARS:
        raise ExtractError("No extractable text.")

    return ExtractedDocument(
        file_type="docx",
        pages=[PageText(page=None, text=text)],
        page_count=None,
    )


def extract_text_file(path: Path, *, file_type: FileType) -> ExtractedDocument:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            text = path.read_text(encoding="latin-1")
        except Exception as exc:  # noqa: BLE001
            raise ExtractError(f"Could not read text file: {exc}") from exc
    except Exception as exc:  # noqa: BLE001
        raise ExtractError(f"Could not read text file: {exc}") from exc

    if len(text.strip()) < MIN_TEXT_CHARS:
        raise ExtractError("No extractable text.")

    return ExtractedDocument(
        file_type=file_type,
        pages=[PageText(page=None, text=text)],
        page_count=None,
    )
